"""Local, source-grounded document investigation for ALG-AI-02."""
from __future__ import annotations
import re
from collections import Counter
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

MAX_FILE_BYTES = 12 * 1024 * 1024
SUPPORTED = {".txt", ".md", ".csv", ".pdf", ".docx"}

@dataclass
class Chunk:
    document: str
    page: int | None
    text: str

class DocumentError(ValueError):
    pass

def extract_file(filename: str, data: bytes) -> list[Chunk]:
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED:
        raise DocumentError(f"Unsupported file type: {ext or 'no extension'}. Use PDF, DOCX, TXT, MD, or CSV.")
    if not data:
        raise DocumentError(f"{filename} is empty.")
    if len(data) > MAX_FILE_BYTES:
        raise DocumentError(f"{filename} exceeds the 12 MB per-file limit.")
    pages: list[tuple[int | None, str]]
    try:
        if ext == ".pdf":
            import fitz
            doc = fitz.open(stream=data, filetype="pdf")
            pages = [(i + 1, p.get_text("text")) for i, p in enumerate(doc)]
            doc.close()
        elif ext == ".docx":
            from docx import Document
            import io
            doc = Document(io.BytesIO(data))
            pages = [(None, "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(" | ".join(c.text for c in row.cells) for t in doc.tables for row in t.rows))]
        else:
            pages = [(None, data.decode("utf-8-sig", errors="replace"))]
    except Exception as exc:
        raise DocumentError(f"Could not read {filename}: {exc}") from exc
    chunks = []
    for page, text in pages:
        # Keep paragraph boundaries useful for citations; split long paragraphs with overlap.
        for para in (p.strip() for p in re.split(r"\n+", text)):
            if not para:
                continue
            while len(para) > 1100:
                cut = para.rfind(" ", 0, 1000)
                cut = cut if cut > 500 else 1000
                chunks.append(Chunk(filename, page, para[:cut]))
                para = para[max(0, cut - 100):]
            chunks.append(Chunk(filename, page, para))
    if not chunks:
        raise DocumentError(f"No selectable text found in {filename}. Scanned PDFs need OCR before upload.")
    return chunks

def _tokens(text: str) -> list[str]:
    return [t for t in re.findall(r"[\w'-]+", text.lower()) if len(t) > 1]

def retrieve(question: str, chunks: Iterable[Chunk], limit: int = 6) -> list[dict]:
    q = _tokens(question)
    if not q:
        raise DocumentError("Enter a question with searchable terms.")
    qset = set(q)
    scored = []
    for ch in chunks:
        toks = _tokens(ch.text)
        if not toks:
            continue
        counts = Counter(toks)
        overlap = sum(1 + __import__('math').log(counts[t]) for t in qset if counts[t])
        # Normalize for document length; this remains transparent keyword retrieval, not an LLM.
        score = overlap / (1 + __import__('math').log(1 + len(toks)))
        if score:
            scored.append((score, ch))
    scored.sort(key=lambda item: item[0], reverse=True)
    best = scored[:limit]
    if not best or best[0][0] < 0.35:
        return []
    max_score = best[0][0]
    return [{"document": ch.document, "page": ch.page, "excerpt": ch.text, "score": round(score / max_score, 3)} for score, ch in best]

def detect_conflicts(question: str, sources: list[dict]) -> list[dict]:
    """Surface opposite-polarity claims across documents when they discuss the same terms."""
    terms = set(_tokens(question))
    claims = []
    for s in sources:
        text = s["excerpt"]
        toks = set(_tokens(text))
        if len(terms & toks) < min(2, len(terms)):
            continue
        low = text.lower()
        neg = bool(re.search(r"\b(no|not|never|cannot|can't|ineligible|contraindicated|prohibited|false|denied|without)\b", low))
        claims.append((s, neg))
    conflicts = []
    for i, (a, neg_a) in enumerate(claims):
        for b, neg_b in claims[i+1:]:
            if neg_a != neg_b and a["document"] != b["document"]:
                conflicts.append({"sources": [a, b], "reason": "These passages use opposite positive/negative language and may conflict. Review the cited context."})
    return conflicts[:3]

def investigate(question: str, chunks: list[Chunk]) -> dict:
    sources = retrieve(question, chunks)
    conflicts = detect_conflicts(question, sources)
    uncertain = not sources or sources[0]["score"] < 0.55
    if not sources:
        answer = "I couldn't find relevant passages in the uploaded documents. Try rephrasing the question or upload a document that covers this topic."
    else:
        answer = "I found potentially relevant passages, but this local demo does not generate a synthesized answer. Review the excerpts below; they are the evidence to use when answering the question."
    return {"question": question, "answer": answer, "sources": sources, "conflicts": conflicts, "uncertainty": {"flagged": uncertain or bool(conflicts), "message": "Evidence is weak or potentially conflicting. The system cannot support a reliable answer from this evidence alone." if uncertain or conflicts else "Retrieved passages are relevant by keyword overlap. Relevance is not proof that a passage answers the question."}}
